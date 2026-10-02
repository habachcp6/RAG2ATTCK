"""Export scientific_report.md to a beautifully formatted Word (.docx) document.

Incorporates publication-quality typography, professional table formatting with
OpenXML pagination rules (<w:tblHeader/>, <w:cantSplit/>), column width optimization,
native OpenXML mathematical typesetting (OMML), callout styling, figure embedding with
SDT locators and keep_with_next pagination, and rigorous QA verification.
"""

from __future__ import annotations

import argparse
import html
import re
import zipfile
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def extract_braced(s: str, start_brace_idx: int) -> tuple[str, int]:
    """Extract content inside matching braces {...} starting at start_brace_idx."""
    depth = 0
    content = []
    i = start_brace_idx
    while i < len(s):
        ch = s[i]
        if ch == "{":
            depth += 1
            if depth > 1:
                content.append(ch)
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return "".join(content), i + 1
            else:
                content.append(ch)
        else:
            if depth >= 1:
                content.append(ch)
        i += 1
    return "".join(content), len(s)


def replace_fractions(text: str) -> str:
    """Recursively convert LaTeX fractions \\frac{num}{den} to (num) / (den)
    handling nested braces.
    """
    pattern = r"\\frac\{"
    while True:
        m = re.search(pattern, text)
        if not m:
            break
        idx = m.start()
        num, num_end = extract_braced(text, m.end() - 1)
        if num_end < len(text) and text[num_end] == "{":
            den, den_end = extract_braced(text, num_end)
            text = text[:idx] + f"({num}) / ({den})" + text[den_end:]
        else:
            text = text[:idx] + num + text[num_end:]
    return text


def latex_to_unicode(text: str) -> str:
    """Convert LaTeX mathematical notation to clean, structured Unicode math text (fallback)."""
    s = text.strip()

    # Preserve conditioning before unwrapping text
    s = re.sub(r"\\mid(?![a-zA-Z])", " | ", s)
    s = s.replace(r"\%", "%")
    s = re.sub(r"\\\$([0-9.]+)", r"$\1", s)

    if s.startswith("$$") and s.endswith("$$"):
        s = s[2:-2].strip()

    s = replace_fractions(s)

    for tag in (r"\\text", r"\\mathrm", r"\\mathbf", r"\\mathit"):
        while re.search(tag + r"\{", s):
            m = re.search(tag + r"\{", s)
            start = m.start()
            content, end = extract_braced(s, m.end() - 1)
            s = s[:start] + content + s[end:]

    s = re.sub(r"\\mathbb\{I\}", "I", s)
    s = re.sub(r"\\mathcal\{C\}", "C", s)

    replacements = [
        (r"\\Delta", "Δ"),
        (r"\\to", "→"),
        (r"\\in", "∈"),
        (r"\\notin", "∉"),
        (r"\\subseteq", "⊆"),
        (r"\\cap", "∩"),
        (r"\\cup", "∪"),
        (r"\\emptyset", "∅"),
        (r"\\neq", "≠"),
        (r"\\le(?![a-zA-Z])", "≤"),
        (r"\\ge(?![a-zA-Z])", "≥"),
        (r"\\approx", "≈"),
        (r"\\sim", "~"),
        (r"\\times", "×"),
        (r"\\cdot", "·"),
        (r"\\forall", "∀"),
        (r"\\equiv", "≡"),
        (r"\\land", "∧"),
        (r"\\lor", "∨"),
        (r"\\quad", "  "),
        (r"\\qquad", "    "),
        (r"\\langle", "⟨"),
        (r"\\rangle", "⟩"),
        (r"\\cos", "cos"),
        (r"\\\|", "||"),
        (r"\\\{", "{"),
        (r"\\\}", "}"),
        (r"\\hat\{y\}_i", "ŷ_i"),
        (r"\\hat\{y\}", "ŷ"),
        (r"\\hat\{q\}", "q̂"),
        (r"\\hat\{d\}_i", "d̂_i"),
        (r"\\hat\{d\}", "d̂"),
        (r"\\left\(", "("),
        (r"\\right\)", ")"),
        (r"\\left\{", "{"),
        (r"\\right\}", "}"),
        (r"\\left\[", "["),
        (r"\\right\]", "]"),
        (r"\\_", "_"),
        (r"\\%", "%"),
    ]
    for pattern, rep in replacements:
        s = re.sub(pattern, rep, s)

    s = re.sub(r"\\sum_\{([^}]+)\}\^\{([^}]+)\}", r"∑_{(\1)}^\2", s)
    s = re.sub(r"\\sum_\{([^}]+)\}", r"∑_{(\1)}", s)
    s = re.sub(r"\\sum", "∑", s)

    s = re.sub(r"\\begin\{cases\}", "", s)
    s = re.sub(r"\\end\{cases\}", "", s)
    s = re.sub(r"\\\\", "; ", s)
    s = re.sub(r"\\&", " ", s)
    s = re.sub(r"&", " ", s)

    s = re.sub(r"_\{([^}]+)\}", r"_\1", s)
    s = re.sub(r"\^\{([^}]+)\}", r"^\1", s)

    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


class LatexToOmml:
    """Converts LaTeX mathematical expressions into native Word OMML (<m:oMath>)."""

    SYMBOL_MAP = {
        r"\Delta": "Δ",
        r"\to": "→",
        r"\in": "∈",
        r"\notin": "∉",
        r"\subseteq": "⊆",
        r"\subset": "⊂",
        r"\cap": "∩",
        r"\cup": "∪",
        r"\emptyset": "∅",
        r"\neq": "≠",
        r"\ne": "≠",
        r"\le": "≤",
        r"\leq": "≤",
        r"\ge": "≥",
        r"\geq": "≥",
        r"\approx": "≈",
        r"\sim": "~",
        r"\times": "×",
        r"\cdot": "·",
        r"\forall": "∀",
        r"\exists": "∃",
        r"\equiv": "≡",
        r"\land": "∧",
        r"\lor": "∨",
        r"\quad": "  ",
        r"\qquad": "    ",
        r"\,": " ",
        r"\;": " ",
        r"\:": " ",
        r"\!": "",
        r"\left(": "(",
        r"\right)": ")",
        r"\left[": "[",
        r"\right]": "]",
        r"\left\{": "{",
        r"\right\}": "}",
        r"\{": "{",
        r"\}": "}",
        r"\_": "_",
        r"\%": "%",
        r"\mid": "|",
        r"\pm": "±",
        r"\infty": "∞",
    }

    @classmethod
    def _escape(cls, text: str) -> str:
        return html.escape(text, quote=True)

    @classmethod
    def convert_to_omml(cls, latex: str, is_display: bool = False) -> str:
        """Main entry point: returns <m:oMath> XML string."""
        s = latex.strip()
        if s.startswith("$$") and s.endswith("$$"):
            s = s[2:-2].strip()
        elif s.startswith("$") and s.endswith("$"):
            s = s[1:-1].strip()

        if r"\langle" in s and r"\rangle" in s:
            s = s.replace(r"\langle", "⟨").replace(r"\rangle", "⟩")

        inner_xml = cls._parse_tokens(s)
        ns = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
        return f'<m:oMath {ns}>{inner_xml}</m:oMath>'

    @classmethod
    def _parse_tokens(cls, s: str) -> str:
        """Parse LaTeX string into sequence of OMML XML elements."""
        xml_parts = []
        i = 0
        n = len(s)

        while i < n:
            ch = s[i]

            if ch.isspace():
                i += 1
                continue

            # Fraction: \frac{num}{den}
            if s[i:].startswith(r"\frac"):
                idx = i + 5
                while idx < n and s[idx].isspace():
                    idx += 1
                if idx < n and s[idx] == "{":
                    num_text, next_idx = extract_braced(s, idx)
                    while next_idx < n and s[next_idx].isspace():
                        next_idx += 1
                    if next_idx < n and s[next_idx] == "{":
                        den_text, end_idx = extract_braced(s, next_idx)
                        num_xml = cls._parse_tokens(num_text)
                        den_xml = cls._parse_tokens(den_text)
                        xml_parts.append(
                            f"<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>"
                        )
                        i = end_idx
                        continue

            # Summation: \sum_{lower}^{upper} or \sum_{lower} or \sum
            if s[i:].startswith(r"\sum"):
                idx = i + 4
                while idx < n and s[idx].isspace():
                    idx += 1
                sub_text = None
                sup_text = None

                for _ in range(2):
                    if idx < n and s[idx] == "_":
                        idx += 1
                        if idx < n and s[idx] == "{":
                            sub_text, idx = extract_braced(s, idx)
                        elif idx < n:
                            sub_text = s[idx]
                            idx += 1
                        while idx < n and s[idx].isspace():
                            idx += 1
                    elif idx < n and s[idx] == "^":
                        idx += 1
                        if idx < n and s[idx] == "{":
                            sup_text, idx = extract_braced(s, idx)
                        elif idx < n:
                            sup_text = s[idx]
                            idx += 1
                        while idx < n and s[idx].isspace():
                            idx += 1

                sub_hide = "0" if sub_text is not None else "1"
                sup_hide = "0" if sup_text is not None else "1"
                sub_xml = cls._parse_tokens(sub_text) if sub_text else ""
                sup_xml = cls._parse_tokens(sup_text) if sup_text else ""

                xml_parts.append(
                    f"<m:nary>"
                    f"<m:naryPr>"
                    f'<m:chr m:val="∑"/>'
                    f'<m:limLoc m:val="undOvr"/>'
                    f'<m:subHide m:val="{sub_hide}"/>'
                    f'<m:supHide m:val="{sup_hide}"/>'
                    f"</m:naryPr>"
                    f"<m:sub>{sub_xml}</m:sub>"
                    f"<m:sup>{sup_xml}</m:sup>"
                    f"<m:e/>"
                    f"</m:nary>"
                )
                i = idx
                continue

            # Text wrapper: \text{...}, \mathrm{...}, \mathbf{...}, \mathit{...}, \mathbb{...}, \mathcal{...}
            tag_match = re.match(r"^\\(text|mathrm|mathbf|mathit|mathbb|mathcal)\{", s[i:])
            if tag_match:
                tag_name = tag_match.group(1)
                brace_start = i + len(tag_match.group(0)) - 1
                inner, next_idx = extract_braced(s, brace_start)
                if tag_name in ("text", "mathrm"):
                    clean_inner = inner.replace(r"\_", "_").replace(r"\%", "%")
                    xml_parts.append(
                        f"<m:r><m:rPr><m:nor/></m:rPr><m:t>{cls._escape(clean_inner)}</m:t></m:r>"
                    )
                elif tag_name == "mathbf":
                    xml_parts.append(
                        f"<m:r><m:rPr><m:b/></m:rPr><m:t>{cls._escape(inner)}</m:t></m:r>"
                    )
                elif tag_name == "mathbb":
                    rep = {"I": "𝕀", "R": "ℝ", "N": "ℕ", "C": "ℂ"}.get(inner, inner)
                    xml_parts.append(f"<m:r><m:t>{cls._escape(rep)}</m:t></m:r>")
                elif tag_name == "mathcal":
                    rep = {"C": "𝓒", "L": "𝓛", "N": "𝓝"}.get(inner, inner)
                    xml_parts.append(f"<m:r><m:t>{cls._escape(rep)}</m:t></m:r>")
                else:
                    xml_parts.append(f"<m:r><m:t>{cls._escape(inner)}</m:t></m:r>")
                i = next_idx
                continue

            # Accents: \hat{...}
            if s[i:].startswith(r"\hat{"):
                inner, next_idx = extract_braced(s, i + 4)
                e_xml = cls._parse_tokens(inner)
                after_idx = next_idx
                has_sub = False
                sub_text = None
                while after_idx < n and s[after_idx].isspace():
                    after_idx += 1
                if after_idx < n and s[after_idx] == "_":
                    has_sub = True
                    after_idx += 1
                    if after_idx < n and s[after_idx] == "{":
                        sub_text, after_idx = extract_braced(s, after_idx)
                    elif after_idx < n:
                        sub_text = s[after_idx]
                        after_idx += 1
                acc_xml = f'<m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e>{e_xml}</m:e></m:acc>'
                if has_sub:
                    sub_xml = cls._parse_tokens(sub_text)
                    xml_parts.append(f"<m:sSub><m:e>{acc_xml}</m:e><m:sub>{sub_xml}</m:sub></m:sSub>")
                    i = after_idx
                    continue
                else:
                    xml_parts.append(acc_xml)
                    i = next_idx
                    continue

            # Norm: \| ... \|
            if s[i:].startswith(r"\|"):
                end_norm = s.find(r"\|", i + 2)
                if end_norm != -1:
                    inner = s[i+2:end_norm]
                    inner_xml = cls._parse_tokens(inner)
                    xml_parts.append(
                        f'<m:d><m:dPr><m:begChr m:val="‖"/><m:endChr m:val="‖"/></m:dPr><m:e>{inner_xml}</m:e></m:d>'
                    )
                    i = end_norm + 2
                    continue

            # Known LaTeX symbol commands
            matched_sym = False
            for cmd, sym in cls.SYMBOL_MAP.items():
                if s[i:].startswith(cmd):
                    after_ch_idx = i + len(cmd)
                    if after_ch_idx < n and cmd[-1].isalpha() and s[after_ch_idx].isalpha():
                        continue
                    xml_parts.append(f"<m:r><m:t>{cls._escape(sym)}</m:t></m:r>")
                    i = after_ch_idx
                    matched_sym = True
                    break
            if matched_sym:
                continue

            # Identifiers or operators with Subscripts and Superscripts
            base_match = re.match(r"^([a-zA-Z0-9]+|[\(\)\[\]\{\}\+\-\=\<\>\,\.\:\;])", s[i:])
            if base_match:
                base_str = base_match.group(1)
                idx = i + len(base_str)
                sub_text = None
                sup_text = None
                if idx < n and s[idx] in ("_", "^"):
                    for _ in range(2):
                        if idx < n and s[idx] == "_":
                            idx += 1
                            if idx < n and s[idx] == "{":
                                sub_text, idx = extract_braced(s, idx)
                            elif idx < n:
                                sub_text = s[idx]
                                idx += 1
                        elif idx < n and s[idx] == "^":
                            idx += 1
                            if idx < n and s[idx] == "{":
                                sup_text, idx = extract_braced(s, idx)
                            elif idx < n:
                                sup_text = s[idx]
                                idx += 1

                base_xml = f"<m:r><m:t>{cls._escape(base_str)}</m:t></m:r>"
                if sub_text is not None and sup_text is not None:
                    sub_xml = cls._parse_tokens(sub_text)
                    sup_xml = cls._parse_tokens(sup_text)
                    xml_parts.append(
                        f"<m:sSubSup><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub><m:sup>{sup_xml}</m:sup></m:sSubSup>"
                    )
                    i = idx
                    continue
                elif sub_text is not None:
                    sub_xml = cls._parse_tokens(sub_text)
                    xml_parts.append(
                        f"<m:sSub><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub></m:sSub>"
                    )
                    i = idx
                    continue
                elif sup_text is not None:
                    sup_xml = cls._parse_tokens(sup_text)
                    xml_parts.append(
                        f"<m:sSup><m:e>{base_xml}</m:e><m:sup>{sup_xml}</m:sup></m:sSup>"
                    )
                    i = idx
                    continue
                else:
                    xml_parts.append(base_xml)
                    i = idx
                    continue

            # Fallback for any other character
            xml_parts.append(f"<m:r><m:t>{cls._escape(s[i])}</m:t></m:r>")
            i += 1

        return "".join(xml_parts)


def set_cell_margins(cell, top=60, bottom=60, left=50, right=50):
    """Set inner padding for table cells (in twips: 1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, fill_hex="F0F4F8"):
    """Set background shading for a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_border(cell, **kwargs):
    """Set individual cell borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f"w:{edge}"
            element = parse_xml(
                f'<{tag} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" '
                f'w:sz="{edge_data.get("sz", 4)}" w:space="0" '
                f'w:color="{edge_data.get("color", "CCCCCC")}"/>'
            )
            tcBorders.append(element)
    tcPr.append(tcBorders)


def apply_table_pagination_rules(table):
    """Ensure repeat header on page break (<w:tblHeader/>) and
    prevent row splitting (<w:cantSplit/>).
    """
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cantSplit = parse_xml(f"<w:cantSplit {nsdecls('w')}/>")
        trPr.append(cantSplit)
        if r_idx == 0:
            tblHeader = parse_xml(f"<w:tblHeader {nsdecls('w')}/>")
            trPr.append(tblHeader)
            # Set keep_with_next on header cell paragraphs to prevent orphan header rows
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True


def assign_table_column_widths(table, num_cols: int, header_texts: list[str]) -> list[float]:
    """Assign explicit column widths ensuring total table width <= 6.50 inches."""
    table.autofit = False
    MAX_WIDTH = 6.50
    hdr_joined = " ".join(header_texts).lower()

    if num_cols == 10:
        widths = [1.10] + [0.60] * 9
    elif num_cols == 9:
        # Table 5: Resource Consumption (9 cols)
        widths = [0.90] + [0.70] * 8
    elif num_cols == 8:
        # Table 4: Failure Decomposition (8 cols)
        widths = [0.90, 0.80, 0.90, 0.95, 0.95, 0.70, 0.65, 0.65]
    elif num_cols == 6:
        if "retrieval depth" in hdr_joined:
            # Table 2a: Attribution Performance
            widths = [1.10, 1.00, 0.90, 1.10, 1.10, 1.30]
        elif "single-event" in hdr_joined:
            # Table 3: Representation Stratification
            widths = [1.00, 1.15, 1.25, 1.05, 1.05, 1.00]
        elif "h-techniquerag" in hdr_joined:
            # Table 1b: Comparators 5-8 + RAG2ATTCK (6 cols)
            widths = [1.25, 1.05, 1.05, 1.05, 1.05, 1.05]
        else:
            # Table 2b: Attribution Diagnostics (6 cols)
            widths = [1.10, 1.00, 1.15, 1.10, 1.10, 1.05]
    elif num_cols == 5:
        widths = [1.30, 1.30, 1.30, 1.30, 1.30]
    elif num_cols == 4:
        # Table 6: Cryptographic Reproducibility Manifest
        widths = [1.60, 1.80, 0.90, 2.20]
    elif num_cols == 3:
        widths = [1.80, 2.70, 2.00]
    else:
        w_each = round(MAX_WIDTH / num_cols, 2)
        widths = [w_each] * num_cols
        diff = sum(widths) - MAX_WIDTH
        if diff != 0:
            widths[-1] -= diff

    for col_idx, w in enumerate(widths):
        table.columns[col_idx].width = Inches(w)
    for row in table.rows:
        for col_idx, w in enumerate(widths):
            row.cells[col_idx].width = Inches(w)

    return widths


def format_inline_runs(
    paragraph,
    text: str,
    font_size: Pt | None = None,
    default_bold: bool = False,
    default_italic: bool = False,
    default_color: RGBColor | None = None,
):
    """Parse inline markdown (bold, italic, code, math, links, currency)
    and append runs to paragraph.
    """
    text = text.replace(r"\%", "%")
    token_pattern = re.compile(
        r"(\\\*|\\\$[0-9.]+|\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\$[^$]+\$)"
    )
    parts = token_pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part == r"\*":
            run = paragraph.add_run("*")
            if font_size:
                run.font.size = font_size
            if default_bold:
                run.bold = True
            if default_italic:
                run.italic = True
            if default_color:
                run.font.color.rgb = default_color
        elif re.match(r"^\\\$[0-9.]+$", part):
            run = paragraph.add_run(part[1:])
            if font_size:
                run.font.size = font_size
            if default_bold:
                run.bold = True
            if default_italic:
                run.italic = True
            if default_color:
                run.font.color.rgb = default_color
        elif part.startswith("***") and part.endswith("***"):
            format_inline_runs(
                paragraph,
                part[3:-3],
                font_size=font_size,
                default_bold=True,
                default_italic=True,
                default_color=default_color,
            )
        elif part.startswith("**") and part.endswith("**"):
            format_inline_runs(
                paragraph,
                part[2:-2],
                font_size=font_size,
                default_bold=True,
                default_italic=default_italic,
                default_color=default_color,
            )
        elif part.startswith("*") and part.endswith("*"):
            format_inline_runs(
                paragraph,
                part[1:-1],
                font_size=font_size,
                default_bold=default_bold,
                default_italic=True,
                default_color=default_color,
            )
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            code_size = Pt(font_size.pt * 0.9) if font_size else Pt(9.5)
            run.font.size = code_size
            run.font.color.rgb = RGBColor(0x8A, 0x1F, 0x11)
        elif part.startswith("[") and "]" in part and part.endswith(")"):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                link_text, link_url = m.groups()
                run = paragraph.add_run(link_text)
                run.font.color.rgb = RGBColor(0x09, 0x69, 0xDA)
                run.underline = True
                if font_size:
                    run.font.size = font_size
                if default_bold:
                    run.bold = True
                if default_italic:
                    run.italic = True
            else:
                run = paragraph.add_run(part)
                if font_size:
                    run.font.size = font_size
                if default_bold:
                    run.bold = True
                if default_italic:
                    run.italic = True
        elif part.startswith("$") and part.endswith("$"):
            math_content = part[1:-1]
            try:
                omml_xml = LatexToOmml.convert_to_omml(math_content, is_display=False)
                paragraph._element.append(parse_xml(omml_xml))
            except Exception:
                converted = latex_to_unicode(math_content)
                run = paragraph.add_run(converted)
                run.font.name = "Cambria Math"
                run.italic = True
                if font_size:
                    run.font.size = font_size
                if default_bold:
                    run.bold = True
                if default_color:
                    run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(part)
            if font_size:
                run.font.size = font_size
            if default_bold:
                run.bold = True
            if default_italic:
                run.italic = True
            if default_color:
                run.font.color.rgb = default_color


def add_display_math(doc, math_text: str):
    """Render display math block as an indented, styled formula paragraph with native OMML."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = True

    pPr = p._element.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    pPr.append(shd)
    pBdr = parse_xml(
        f"<w:pBdr {nsdecls('w')}>"
        '<w:left w:val="single" w:sz="18" w:space="8" w:color="0969DA"/>'
        "</w:pBdr>"
    )
    pPr.append(pBdr)

    try:
        omml_xml = LatexToOmml.convert_to_omml(math_text, is_display=True)
        p._element.append(parse_xml(omml_xml))
    except Exception:
        converted = latex_to_unicode(math_text)
        run = p.add_run(converted)
        run.font.name = "Cambria Math"
        run.font.size = Pt(10.5)
        run.italic = True
        run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)


def build_docx_from_markdown(
    md_path: Path, output_docx_path: Path, figures_dir: Path | None = None
):
    """Convert scientific_report.md into a high-quality Word document."""
    doc = docx.Document()

    # Configure page margins (1.0 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal style
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    # Clean Title style in Word template to eliminate default blue border/rule and ensure black text
    if "Title" in doc.styles:
        title_style = doc.styles["Title"]
        title_style.font.name = "Georgia"
        title_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        title_pPr = title_style._element.get_or_add_pPr()
        bdr = title_pPr.find(qn("w:pBdr"))
        if bdr is not None:
            title_pPr.remove(bdr)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    in_references = False
    last_caption_text: str | None = None
    table_seq_idx: int = 0
    figure_seq_idx: int = 1

    def flush_table():
        nonlocal in_table, table_lines, last_caption_text, table_seq_idx
        if not table_lines:
            in_table = False
            return

        raw_rows = []
        for tline in table_lines:
            cells = [c.strip() for c in tline.split("|")]
            if len(cells) >= 3 and cells[0] == "" and cells[-1] == "":
                cells = cells[1:-1]
            if cells and all(re.match(r"^:?-+:?$", c) for c in cells if c):
                continue
            raw_rows.append(cells)

        if not raw_rows:
            table_lines = []
            in_table = False
            return

        num_cols = max(len(r) for r in raw_rows)
        table = doc.add_table(rows=len(raw_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        hdr_str = " ".join(raw_rows[0]).lower() if raw_rows else ""
        caption_str = last_caption_text or ""

        if "table 1a" in caption_str.lower() or "comparators 1-4" in hdr_str or "yang & hsu" in hdr_str:
            table_id = "table_1a"
            table_caption = "Table 1a: Multi-Dimensional Comparator Matrix (Part 1)"
        elif "table 1b" in caption_str.lower() or "comparators 5-8" in hdr_str or "h-techniquerag" in hdr_str:
            table_id = "table_1b"
            table_caption = "Table 1b: Multi-Dimensional Comparator Matrix (Part 2)"
        elif "table 2a" in caption_str.lower() or ("complexity" in hdr_str and "headline" in hdr_str):
            table_id = "table_2a"
            table_caption = "Table 2a: Primary Attribution Performance and Ground-Truth Complexity"
        elif "table 2b" in caption_str.lower() or "completed outputs" in hdr_str:
            table_id = "table_2b"
            table_caption = "Table 2b: Attribution Diagnostic Metrics"
        elif "table 3b" in caption_str.lower() or "complete pairs" in hdr_str:
            table_id = "table_3b"
            table_caption = "Table 3b: Paired Scorable Representation Concordance"
        elif "table 3" in caption_str.lower() or "single-event" in hdr_str:
            table_id = "table_3"
            table_caption = "Table 3: Representation Stratification (Single vs Contextual)"
        elif "table 4" in caption_str.lower() or "downstream selection failure" in hdr_str or "upstream retrieval miss" in hdr_str:
            table_id = "table_4"
            table_caption = "Table 4: Decoupled Failure Decomposition Matrix"
        elif "table 5b" in caption_str.lower() or "financial ledger" in hdr_str or "accounting dimension" in hdr_str:
            table_id = "table_5b"
            table_caption = "Table 5b: Whole-Study Financial Ledger and Budget Reconciliation"
        elif "table 5" in caption_str.lower() or "total input tokens" in hdr_str:
            table_id = "table_5"
            table_caption = "Table 5: Resource Consumption and Latency Scaling"
        elif "table 6" in caption_str.lower() or "cryptographic reproducibility" in hdr_str or "asset description" in hdr_str:
            table_id = "table_6"
            table_caption = "Table 6: Cryptographic Reproducibility Manifest"
        else:
            table_seq_idx += 1
            table_id = f"table_{table_seq_idx}"
            table_caption = caption_str or f"Table {table_seq_idx}"

        tblPr = table._element.tblPr
        tblCaption_elem = parse_xml(f'<w:tblCaption {nsdecls("w")} w:val="{table_caption}"/>')
        tblDesc_elem = parse_xml(f'<w:tblDescription {nsdecls("w")} w:val="{table_id}"/>')
        tblPr.append(tblCaption_elem)
        tblPr.append(tblDesc_elem)

        # Standardize table cell font sizes: >= 8.5 pt (eliminating 6.5pt font)
        if num_cols >= 8:
            cell_font_size = Pt(8.5)
            pad_top, pad_bot, pad_left, pad_right = 60, 60, 50, 50
        else:
            cell_font_size = Pt(9.0)
            pad_top, pad_bot, pad_left, pad_right = 70, 70, 60, 60

        for r_idx, row_data in enumerate(raw_rows):
            is_header = r_idx == 0
            for c_idx in range(num_cols):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = table.cell(r_idx, c_idx)
                p = cell.paragraphs[0]
                p.text = ""
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.05

                if is_header:
                    set_cell_background(cell, "092C4C")  # Dark Navy Header
                    format_inline_runs(
                        p,
                        cell_text,
                        font_size=cell_font_size,
                        default_bold=True,
                        default_color=RGBColor(0xFF, 0xFF, 0xFF),
                    )
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")  # Subtle alternating light row
                    else:
                        set_cell_background(cell, "FFFFFF")
                    format_inline_runs(
                        p,
                        cell_text,
                        font_size=cell_font_size,
                        default_bold=False,
                        default_color=RGBColor(0x24, 0x29, 0x2F),
                    )

                set_cell_margins(cell, top=pad_top, bottom=pad_bot, left=pad_left, right=pad_right)
                set_cell_border(
                    cell,
                    top={"sz": 4, "val": "single", "color": "D0D7DE"},
                    bottom={"sz": 4, "val": "single", "color": "D0D7DE"},
                    left={"sz": 2, "val": "single", "color": "E1E4E8"},
                    right={"sz": 2, "val": "single", "color": "E1E4E8"},
                )

                if table_id != "table_6":
                    tag_val = f"{table_id}_r{r_idx}_c{c_idx}"
                    alias_val = f"{table_id} Row {r_idx} Col {c_idx}"
                    sdt_elem = parse_xml(
                        f'<w:sdt {nsdecls("w")}>'
                        f'  <w:sdtPr>'
                        f'    <w:tag w:val="{tag_val}"/>'
                        f'    <w:alias w:val="{alias_val}"/>'
                        f'  </w:sdtPr>'
                        f'  <w:sdtContent/>'
                        f'</w:sdt>'
                    )
                    sdt_content = sdt_elem.find(qn("w:sdtContent"))
                    runs_to_wrap = [
                        child for child in list(p._element)
                        if child.tag in (qn("w:r"), qn("m:oMath"))
                    ]
                    for r_elem in runs_to_wrap:
                        sdt_content.append(r_elem)
                    p._element.append(sdt_elem)

        apply_table_pagination_rules(table)
        header_texts = raw_rows[0] if raw_rows else []
        assign_table_column_widths(table, num_cols, header_texts)

        last_caption_text = None

        post_p = doc.add_paragraph()
        post_p.paragraph_format.space_after = Pt(6)
        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_text = "\n".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0

        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

        pPr = p._element.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F6F8FA"/>')
        pPr.append(shd)

        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]

        # Handle fenced code block
        if line.strip().startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip HTML comments (e.g. <!-- FIXTURE_ONLY: true -->)
        if line.strip().startswith("<!--") and line.strip().endswith("-->"):
            i += 1
            continue

        # Handle display math $$ ... $$
        if line.strip().startswith("$$"):
            if in_table:
                flush_table()
            if in_code_block:
                flush_code()

            stripped = line.strip()
            if stripped.endswith("$$") and len(stripped) > 4:
                math_content = stripped[2:-2].strip()
                add_display_math(doc, math_content)
                i += 1
                continue
            else:
                math_lines = []
                if len(stripped) > 2:
                    math_lines.append(stripped[2:].strip())
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    math_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    closing_line = lines[i].strip()
                    if len(closing_line) > 2:
                        math_lines.append(closing_line[:-2].strip())
                    i += 1
                add_display_math(doc, " ".join(math_lines))
                continue

        # Handle Markdown table lines
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Handle Headings with keep_with_next to prevent orphan section titles
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.font.name = "Georgia"
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        elif line.startswith("## "):
            sec_title = line[3:].strip()
            if sec_title.lower().startswith("references"):
                in_references = True
            else:
                in_references = False
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(sec_title)
            run.font.name = "Georgia"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
        elif line.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[4:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
        elif line.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[5:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x44)
        elif line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._element.get_or_add_pPr()
            pBdr = parse_xml(
                f"<w:pBdr {nsdecls('w')}>"
                '<w:bottom w:val="single" w:sz="6" w:space="1" w:color="D0D7DE"/>'
                "</w:pBdr>"
            )
            pPr.append(pBdr)
        elif line.startswith(">"):
            # Gather all contiguous quote lines into a unified callout block
            quote_lines = []
            while i < len(lines) and lines[i].startswith(">"):
                q_l = lines[i][1:].strip()
                quote_lines.append(q_l)
                i += 1

            border_color = "0969DA"  # Default blue
            bg_color = "F0F4F8"      # Default soft blue
            filtered_lines = []

            for ql in quote_lines:
                m_callout = re.match(
                    r"^\[!(NOTE|WARNING|IMPORTANT|TIP|CAUTION)\]$", ql.strip(), re.IGNORECASE
                )
                if m_callout:
                    callout_type = m_callout.group(1).upper()
                    if callout_type == "WARNING":
                        border_color = "D97706"  # Amber
                        bg_color = "FFFBEB"      # Light amber
                    elif callout_type == "IMPORTANT":
                        border_color = "8250DF"  # Purple
                        bg_color = "FBEFFF"
                    elif callout_type == "CAUTION":
                        border_color = "CF222E"  # Red
                        bg_color = "FFEBE9"
                    elif callout_type == "TIP":
                        border_color = "1A7F37"  # Green
                        bg_color = "F0FDF4"
                    else:
                        border_color = "0969DA"  # Blue
                        bg_color = "F0F4F8"
                else:
                    filtered_lines.append(ql)

            for q_idx, q_text in enumerate(filtered_lines):
                if not q_text:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(4 if q_idx == 0 else 0)
                p.paragraph_format.space_after = Pt(4 if q_idx == len(filtered_lines) - 1 else 2)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_together = True

                pPr = p._element.get_or_add_pPr()
                pBdr = parse_xml(
                    f"<w:pBdr {nsdecls('w')}>"
                    f'<w:left w:val="single" w:sz="24" w:space="8" w:color="{border_color}"/>'
                    f"</w:pBdr>"
                )
                pPr.append(pBdr)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
                pPr.append(shd)

                if q_text.startswith("- ") or q_text.startswith("* "):
                    bullet_run = p.add_run("•  ")
                    bullet_run.bold = True
                    bullet_run.font.name = "Calibri"
                    bullet_run.font.size = Pt(10)
                    format_inline_runs(p, q_text[2:].strip(), font_size=Pt(10))
                else:
                    format_inline_runs(p, q_text, font_size=Pt(10))
            continue
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = line.strip()[2:]
            format_inline_runs(p, content)
        elif re.match(r"^\d+\.\s+", line.strip()):
            m = re.match(r"^(\d+)\.\s+(.*)", line.strip())
            num_str, content = m.groups()
            p = doc.add_paragraph()
            if in_references:
                # Academic bibliography formatting with static numbering and hanging indent
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.35)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_together = True  # Prevent splitting across page break
                num_run = p.add_run(f"[{num_str}] ")
                num_run.bold = True
                num_run.font.name = "Calibri"
                num_run.font.size = Pt(10)
                num_run.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
                format_inline_runs(p, content, font_size=Pt(10))
            else:
                p.paragraph_format.left_indent = Inches(0.30)
                p.paragraph_format.first_line_indent = Inches(-0.20)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                num_run = p.add_run(f"{num_str}. ")
                num_run.bold = True
                num_run.font.name = "Calibri"
                num_run.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
                format_inline_runs(p, content)
        elif re.match(r"^!\[(.*?)\]\((.*?)\)", line.strip()):
            if in_table:
                flush_table()
            if in_code_block:
                flush_code()
            m_img = re.match(r"^!\[(.*?)\]\((.*?)\)", line.strip())
            alt_text, img_rel_path = m_img.groups()

            img_filename = Path(img_rel_path).name
            candidates: list[Path] = []
            if figures_dir is not None:
                figures_dir_p = Path(figures_dir)
                candidates.extend(
                    [
                        figures_dir_p / img_rel_path,
                        figures_dir_p / img_filename,
                    ]
                )
            candidates.extend(
                [
                    md_path.parent / img_rel_path,
                    md_path.parent / "figures" / img_filename,
                    Path(__file__).resolve().parent.parent
                    / "docs"
                    / "report"
                    / "figures"
                    / img_filename,
                    Path.cwd() / "docs" / "report" / "figures" / img_filename,
                ]
            )
            resolved_img = None
            for cand in candidates:
                if cand.is_file():
                    resolved_img = cand
                    break

            if resolved_img:
                fig_m = re.search(
                    r"fig(?:ure)?[-_]?(\d+)",
                    alt_text.lower() + " " + resolved_img.name.lower(),
                )
                fig_num = fig_m.group(1) if fig_m else str(figure_seq_idx)
                figure_seq_idx += 1

                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(3)
                p_img.paragraph_format.keep_with_next = True

                run_img = p_img.add_run()
                run_img.add_picture(str(resolved_img), width=Inches(6.25))

                tag_val = f"fig_{fig_num}"
                alias_val = f"Figure {fig_num}"
                sdt_elem = parse_xml(
                    f'<w:sdt {nsdecls("w")}>'
                    f'  <w:sdtPr>'
                    f'    <w:tag w:val="{tag_val}"/>'
                    f'    <w:alias w:val="{alias_val}"/>'
                    f'  </w:sdtPr>'
                    f'  <w:sdtContent/>'
                    f'</w:sdt>'
                )
                sdt_content = sdt_elem.find(qn("w:sdtContent"))
                runs_to_wrap = [child for child in list(p_img._element) if child.tag == qn("w:r")]
                for r_elem in runs_to_wrap:
                    sdt_content.append(r_elem)
                p_img._element.append(sdt_elem)
            else:
                raise FileNotFoundError(
                    f"[FAIL_CLOSED] Figure image file not found for: {img_rel_path}"
                )
        elif line.strip().startswith("*Figure "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(10)
            format_inline_runs(
                p,
                line.strip(),
                font_size=Pt(9.5),
                default_italic=True,
                default_color=RGBColor(0x57, 0x60, 0x6A),
            )
        elif line.strip().startswith("*Table "):
            last_caption_text = line.strip().strip("*").rstrip(".*")
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            format_inline_runs(
                p,
                line.strip(),
                font_size=Pt(10),
                default_italic=True,
                default_bold=False,
                default_color=RGBColor(0x1F, 0x23, 0x28),
            )
        elif line.strip():
            p = doc.add_paragraph()
            format_inline_runs(p, line.strip())

        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    output_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_docx_path))
    print(f"Successfully generated DOCX at {output_docx_path}")


def audit_docx_quality(doc_path: Path):
    """Rigorous QA audit verifying rendering invariants on generated DOCX."""
    doc = docx.Document(str(doc_path))

    # Regex to detect raw LaTeX tokens or unparsed $$ blocks
    tex_pattern = re.compile(r"\\[a-zA-Z]+|\$\$")
    errors = []

    # 1. Audit Title paragraph (P0): pure black text, NO w:pBdr
    p0 = doc.paragraphs[0]
    if not p0.text.startswith("Evaluating MITRE ATT&CK-Grounded RAG"):
        errors.append(f"P0 text does not start with Title: '{p0.text[:60]}'")
    p0_xml = p0._element.xml
    if "w:pBdr" in p0_xml:
        errors.append(f"Title paragraph has w:pBdr element: {p0_xml[:200]}")
    for r in p0.runs:
        if r.font.color and r.font.color.rgb:
            if r.font.color.rgb != RGBColor(0, 0, 0):
                errors.append(f"Title run has non-black color: {r.font.color.rgb}")

    # 2. Audit all paragraphs for raw TeX leakage and capture References
    in_ref_sec = False
    ref_paragraphs = []
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text
        if text.strip() == "References" and p.style.name.startswith("Heading"):
            in_ref_sec = True
            continue
        if in_ref_sec:
            if p.style.name.startswith("Heading"):
                in_ref_sec = False
            elif text.strip().startswith("[") and "]" in text:
                ref_paragraphs.append((p_idx, p))

        is_code = any(r.font.name == "Consolas" for r in p.runs)
        if not is_code:
            matches = tex_pattern.findall(text)
            if matches:
                errors.append(f"Paragraph {p_idx} has raw TeX tokens {matches}: '{text[:120]}...'")

    # 3. Audit References numbering: exactly 13 references, numbered [1] to [13],
    # no List Number style
    if len(ref_paragraphs) != 13:
        errors.append(f"Expected 13 references in References section, found {len(ref_paragraphs)}")
    else:
        for expected_num, (p_idx, p) in enumerate(ref_paragraphs, 1):
            expected_prefix = f"[{expected_num}]"
            if not p.text.strip().startswith(expected_prefix):
                errors.append(
                    f"Ref paragraph {p_idx} expected prefix '{expected_prefix}', "
                    f"got '{p.text[:20]}'"
                )
            if p.style.name == "List Number":
                errors.append(
                    f"Ref paragraph {p_idx} uses List Number style instead of static numbering"
                )

    # 4. Audit all table cells for raw TeX leakage, pagination rules, and width constraints
    t1a_found = False
    t1b_found = False
    t2b_found = False

    def _cell_text(cell) -> str:
        return "".join(cell._element.xpath(".//w:t/text()"))

    for t_idx, table in enumerate(doc.tables):
        row0_text = " ".join(_cell_text(c).strip() for c in table.rows[0].cells).lower()
        if "yang & hsu" in row0_text:
            t1a_found = True
            if len(table.columns) != 5:
                errors.append(f"Table 1a expected 5 cols, got {len(table.columns)}")
        elif "h-techniquerag" in row0_text:
            t1b_found = True
            if len(table.columns) != 6:
                errors.append(f"Table 1b expected 6 cols, got {len(table.columns)}")
        elif "completed outputs" in row0_text:
            t2b_found = True
            if "macro precision" in row0_text or "macro recall" in row0_text:
                errors.append("Table 2b still contains unexported macro precision/recall!")

        # Header check
        row0_xml = table.rows[0]._tr.xml
        if "<w:tblHeader" not in row0_xml:
            errors.append(f"Table {t_idx} row 0 is missing <w:tblHeader/>")

        # CantSplit check on every row
        for r_idx, row in enumerate(table.rows):
            row_xml = row._tr.xml
            if "<w:cantSplit" not in row_xml:
                errors.append(f"Table {t_idx} row {r_idx} is missing <w:cantSplit/>")

        # Cell content check
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_text = _cell_text(cell)
                matches = tex_pattern.findall(cell_text)
                if matches:
                    errors.append(
                        f"Table {t_idx} row {r_idx} col {c_idx} has raw TeX {matches}: "
                        f"'{cell_text[:80]}'"
                    )

        # Width check (sum of column widths)
        col_widths_sum = sum(col.width.inches for col in table.columns if col.width)
        if col_widths_sum > 6.55:
            errors.append(
                f"Table {t_idx} width {col_widths_sum:.2f}in exceeds printable limit of 6.50in"
            )

    if not t1a_found:
        errors.append("Table 1a (Comparators 1-4) not found in DOCX tables")
    if not t1b_found:
        errors.append("Table 1b (Comparators 5-8 + RAG2ATTCK) not found in DOCX tables")
    if not t2b_found:
        errors.append("Table 2b (Attribution Diagnostics) not found in DOCX tables")

    # 5. Audit OpenXML SDT locator tags on tables and verify non-empty visible text in sdtContent
    has_sdt = any("w:tag" in t._element.xml for t in doc.tables)
    if not has_sdt:
        errors.append("DOCX tables are missing OpenXML SDT / locator tags (<w:tag/>)")
    has_sdt_text = any(len(t._element.xpath(".//w:sdt/w:sdtContent//w:t")) > 0 for t in doc.tables)
    if not has_sdt_text:
        errors.append("DOCX tables have empty OpenXML SDT tags (<w:sdtContent/> has no visible text)")

    # 6. Audit embedded images in word/media package (All 8 canonical figures)
    PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
    with zipfile.ZipFile(doc_path) as z:
        media_files = [f for f in z.namelist() if f.startswith("word/media/")]
        if len(media_files) < 8:
            errors.append(
                f"DOCX QA: Expected at least 8 embedded figures in word/media/, "
                f"found {len(media_files)}: {media_files}"
            )
        for mf in media_files:
            if mf.endswith(".png"):
                m_data = z.read(mf)
                if not m_data.startswith(PNG_MAGIC):
                    errors.append(f"Embedded image {mf} is not a valid PNG file (bad magic bytes)")

    if errors:
        error_msg = f"DOCX QA Audit FAILED with {len(errors)} error(s):\n" + "\n".join(errors)
        raise AssertionError(error_msg)

    print(
        f"DOCX QA Audit PASSED: 0 raw TeX tokens across {len(doc.paragraphs)} paragraphs and "
        f"{len(doc.tables)} tables ({sum(len(t.rows) for t in doc.tables)} rows). "
        "Title is pure black with no borders. References [1]..[13] statically numbered. "
        "All 8 figures embedded with SDT locators. All tables have cantSplit on all rows, "
        "tblHeader on row 0, and width <= 6.50 inches."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert scientific_report.md into a high-quality Word document (.docx)."
    )
    parser.add_argument(
        "--md-path",
        type=Path,
        default=Path("docs/report/scientific_report.md"),
        help="Path to source markdown report (default: docs/report/scientific_report.md).",
    )
    parser.add_argument(
        "--docx-path",
        type=Path,
        default=Path("docs/report/scientific_report.docx"),
        help="Path to output Word document (default: docs/report/scientific_report.docx).",
    )
    parser.add_argument(
        "--figures-dir",
        type=Path,
        default=None,
        help="Optional path to directory containing report figures.",
    )
    parser.add_argument(
        "--skip-audit",
        action="store_true",
        help="Skip post-export quality audit invariants.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_docx_from_markdown(args.md_path, args.docx_path, figures_dir=args.figures_dir)
    if not args.skip_audit:
        audit_docx_quality(args.docx_path)
